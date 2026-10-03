"""Resolve observed ICAO codes through EASA's official Doc 29 proxy table."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile
import csv
from io import TextIOWrapper

from openpyxl import load_workbook


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def _availability(archive_path: Path) -> tuple[set[str], dict[tuple[str, str], tuple[str, bool, int, int]], dict[tuple[str, str], int]]:
    with ZipFile(archive_path) as z:
        with z.open("ANP2.3_Aircraft.csv") as raw:
            aircraft = list(csv.DictReader(TextIOWrapper(raw, encoding="cp1252", newline=""), delimiter=";"))
        with z.open("ANP2.3_NPD_data.csv") as raw:
            npd = list(csv.DictReader(TextIOWrapper(raw, encoding="cp1252", newline=""), delimiter=";"))
        with z.open("ANP2.3_Default_approach_procedural_steps.csv") as raw:
            arrivals = list(csv.DictReader(TextIOWrapper(raw, encoding="cp1252", newline=""), delimiter=";"))
        with z.open("ANP2.3_Default_departure_procedural_steps.csv") as raw:
            departures = list(csv.DictReader(TextIOWrapper(raw, encoding="cp1252", newline=""), delimiter=";"))
        with z.open("ANP2.3_Default_weights.csv") as raw:
            weights = list(csv.DictReader(TextIOWrapper(raw, encoding="cp1252", newline=""), delimiter=";"))
    by_id = {r["ACFT_ID"].strip(): r for r in aircraft}
    npd_ids = {r["NPD_ID"].strip() for r in npd}
    npd_modes = {(r["NPD_ID"].strip(), r["Noise Metric"].strip(), r["Op Mode"].strip()) for r in npd}
    profile_points = {}
    for op, rows in (("A", arrivals), ("D", departures)):
        for row in rows:
            k=(row["ACFT_ID"].strip(), op)
            profile_points[k]=profile_points.get(k,0)+1
    weight_rows = {}
    for row in weights:
        k=row["ACFT_ID"].strip()
        weight_rows[k]=weight_rows.get(k,0)+1
    results={}
    for acft_id, a in by_id.items():
        npd_id=a["NPD_ID"].strip()
        for operation in ("A","D"):
            results[(acft_id,operation)]=(npd_id, (npd_id,"SEL",operation) in npd_modes,
                                          profile_points.get((acft_id,operation),0), weight_rows.get(acft_id,0))
    return set(by_id), results, profile_points


def run(proximity_path: Path, substitution_path: Path, anp_archive: Path) -> dict:
    proximity=json.loads(proximity_path.read_text())
    local=proximity["airport_proximity_counts"]["within_3000m"]["type_code_path_fragments"]
    wb=load_workbook(substitution_path,read_only=True,data_only=True)
    sheet=wb["by ICAO code"]
    rows=list(sheet.iter_rows(values_only=True))
    headers=[str(x).strip() if x is not None else "" for x in rows[0]]
    by_code={}
    for values in rows[1:]:
        row=dict(zip(headers,values))
        code=str(row.get("ICAO_CODE") or "").strip().upper()
        if code:
            by_code.setdefault(code,[]).append(row)
    anp_ids, npd_avail, profile_points=_availability(anp_archive)
    coverage=[]
    for code, count in sorted(local.items(),key=lambda x:(-x[1],x[0])):
        candidates=by_code.get(code,[])
        if not candidates:
            coverage.append({"icao_type_code":code,"near_vny_path_fragments_not_flight_count":count,
              "substitution_table_rows":0,"status":"no_official_substitution_row_in_2018_jets_heavy_props_table",
              "accepted_for_model":False})
            continue
        operations={}
        invalid=[]
        for op,delta_col in (("D","DELTA_DEP_dB"),("A","DELTA_APP_dB")):
            usable=[]
            for r in candidates:
                proxy=str(r.get("ANP_PROXY") or "").strip()
                raw_delta=r.get(delta_col)
                try: delta=float(raw_delta)
                except (TypeError,ValueError): continue
                if proxy not in anp_ids:
                    invalid.append({"proxy":proxy,"operation":op,"reason":"proxy_absent_from_ANP_v2_3_aircraft_table"}); continue
                npd_id,has_sel,steps,weight_rows=npd_avail[(proxy,op)] if (proxy,op) in npd_avail else (None,False,0,0)
                if not has_sel or steps<2 or weight_rows<1:
                    invalid.append({"proxy":proxy,"operation":op,"reason":"no_complete_SEL_NPD_and_default_procedural_profile_or_weight"}); continue
                usable.append((delta,r,npd_id,steps,weight_rows))
            if usable:
                delta,r,npd_id,steps,weight_rows=max(usable,key=lambda x:(x[0],str(x[1].get("ANP_PROXY"))))
                operations[op]={"selected_ANP_proxy_internal_only":str(r["ANP_PROXY"]).strip(),
                    "delta_db_conservative_max_across_table_variants":delta,
                    "implied_equivalent_movement_factor_for_reference_only":round(10**(delta/10),6),
                    "N_adjustment_applied":False,
                    "why_N_is_not_applied":"The mapped Δ dB adjustment is selected for single-event NPD; Doc 29 treats Δ and equivalent-movement N as alternate representations, so they are not applied together.",
                    "npd_id_internal_only":npd_id,"sel_NPD_available":has_sel,
                    "default_procedural_profile_step_rows":steps,"default_weight_stage_rows":weight_rows,
                    "selected_variant":str(r.get("AIRCRAFT_VARIANT") or "unspecified in code sheet"),
                    "source_engine_manufacturer":r.get("ENGINE_MANUFACTURER"),"source_engine_type":r.get("ENGINE_TYPE"),
                    "delta_source_rows_for_code":len(candidates),"selection_rule":"largest Δ for this operation across all published by-ICAO-code variant rows with a complete ANP proxy SEL table and default procedural profile/weight; conservative per Doc 29 Vol 2 Appendix G5."}
            else:
                operations[op]={"status":"no_usable_proxy_for_operation","delta_source_rows_for_code":len(candidates)}
        ready=all("selected_ANP_proxy_internal_only" in operations.get(op,{}) for op in ("A","D"))
        coverage.append({"icao_type_code":code,"near_vny_path_fragments_not_flight_count":count,
            "substitution_table_rows":len(candidates),"status":"official_conservative_proxy_candidate_pending_root_review" if ready else "incomplete_or_unsupported_mapping",
            "accepted_for_model":False,"operation_candidates":operations,"excluded_candidate_rows":invalid[:20]})
    return {"schema":"quiet_la_vny_doc29_icao_substitution_readiness_internal_v1",
      "activity_source_file":proximity["source_activity_file"],"activity_radius_m":3000,
      "activity_counts_are_path_fragments_not_flights":True,
      "sources":{"substitutions_url":"https://www.easa.europa.eu/en/domains/environment/policy-support-and-research/aircraft-noise-and-performance-anp-data",
        "substitution_download_url":"https://www.easa.europa.eu/en/downloads/138165/en",
        "substitution_file":substitution_path.name,"substitution_sha256":_sha256(substitution_path),
        "published_list_date":"2018-02-22","doc29_method":"ECAC Doc 29 5th Edition Vol. 2 Appendix G5 (pp. G-9–G-10); when only ICAO code is known, choose largest Δ or N across variants or average; this candidate uses maximum Δ for event NPD.",
        "anp_archive_file":anp_archive.name,"anp_archive_sha256":_sha256(anp_archive)},
      "conservative_variant_rule":"For an unknown engine/MTOW variant, choose the greatest per-operation Δ dB among documented ICAO-code rows whose proxy exists in the private ANP v2.3 archive with matching SEL NPD and default fixed-point profile. Use Δ in event NPD, never apply the equivalent N movement factor together. A/D may select different proxies if the table specifies them.",
      "coverage":coverage,
      "unmapped_or_uncertain_activity":"All unlisted codes, including rotorcraft and military/unknown classes, remain unsupported; no blind family proxy is introduced. The output holds internal-only ANP identifiers and must not be redistributed.",
      "model_readiness":"This resolves source-supported proxy candidates only. A local event-Sel calculation still needs accepted movement inference/runway geometry, profile applicability/thrust and altitude datum handling, receiver coordinates/height/terrain, and root approval. No NPD output or measured-series fitting occurs here."}


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--proximity",type=Path,required=True)
    ap.add_argument("--substitutions",type=Path,required=True)
    ap.add_argument("--anp-archive",type=Path,required=True)
    ap.add_argument("--output",type=Path,required=True)
    a=ap.parse_args(); result=run(a.proximity,a.substitutions,a.anp_archive)
    a.output.parent.mkdir(parents=True,exist_ok=True); a.output.write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps({"types":len(result["coverage"]),"path_fragments":sum(r["near_vny_path_fragments_not_flight_count"] for r in result["coverage"]),"official_ICAO_substitution_rows":sum(r["substitution_table_rows"]>0 for r in result["coverage"]),"complete_candidate_types":sum(r["status"]=="official_conservative_proxy_candidate_pending_root_review" for r in result["coverage"]),"unmapped_path_fragments":sum(r["near_vny_path_fragments_not_flight_count"] for r in result["coverage"] if r["substitution_table_rows"]==0)},indent=2))


if __name__=="__main__": main()
