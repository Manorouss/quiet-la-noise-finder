"""Generic internal VNY event adapter for ANP default fixed-point profiles.

Only airport activity windows with required ADS-B witnesses and a qualified
aircraft mapping are evaluated. The fixed-point path builder deliberately does
not consume procedural-step rows. Raw ANP tables stay private; output is internal.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile
from io import TextIOWrapper

from openpyxl import load_workbook

from .anp_loader import load_anp_npd_table
from .doc29 import Doc29InputError, energy_sum_db, npd_level
from .independent_profile_acoustics import _installation_correction, _segment_geometry
from .profile_geometry import segment_profile_points
from .segment_reference import (
    FT_TO_M, KNOT_TO_MPS, _finite_segment_correction, _impedance_adjustment,
    _lateral_attenuation, _start_of_roll_jet, _start_of_roll_turboprop,
)
from .vny_pa28_monitor_site_sensitivity import (
    HEIGHTS_M, VERTICAL_OFFSETS_M, _local_daypart, _local_track_offsets,
    _read_sites, _subset_cnel,
)

MANUAL_SCENARIOS = {
    # Explicitly authorized internal scenario; this is not an ICAO substitution.
    "P28A": {"anp_id": "PA28", "delta_db": 0.0,
             "mapping_kind": "manual_family_variant_scenario_not_exact_match"},
}
OPERATION_CODE = {"arrival": "A", "departure": "D"}
RECEIVER_HEIGHTS_M = HEIGHTS_M
REFERENCE_SPEED_MPS = 160.0 * KNOT_TO_MPS


def _sha(path: Path) -> str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda:f.read(1 << 20),b""): h.update(block)
    return h.hexdigest()


def _csv_rows(z: ZipFile, name: str) -> list[dict]:
    with z.open(name) as raw:
        return list(csv.DictReader(TextIOWrapper(raw,encoding="cp1252",newline=""),delimiter=";"))


def _power_unit(parameter: str) -> str:
    p=parameter.lower()
    if "rpm" in p: return "rpm_per_engine"
    if "%" in p or "percent" in p: return "percent_per_engine"
    if "hp" in p or "horse_power" in p: return "hp_per_engine"
    if "thrust" in p or "cnt" in p or "lb" in p: return "lb_per_engine"
    raise ValueError(f"unsupported ANP power parameter {parameter!r}")


def load_fixed_profile(archive: Path, anp_id: str, operation: str, stage_length: int=1) -> tuple[list[dict],dict]:
    """Load one actual default fixed-point profile, retaining its units and identity."""
    if operation not in ("A","D"): raise ValueError("operation must be A or D")
    with ZipFile(archive) as z:
        aircraft_rows=_csv_rows(z,"ANP2.3_Aircraft.csv")
        fixed_rows=_csv_rows(z,"ANP2.3_Default_fixed_point_profiles.csv")
        weights=_csv_rows(z,"ANP2.3_Default_weights.csv")
    aircraft=[r for r in aircraft_rows if r["ACFT_ID"].strip()==anp_id]
    if len(aircraft)!=1: raise ValueError(f"expected one ANP aircraft row for {anp_id!r}")
    ac=aircraft[0]; unit=_power_unit(ac["Power Parameter"])
    selected=[r for r in fixed_rows if r["ACFT_ID"].strip()==anp_id and r["Op Type"].strip()==operation
              and r["Profile_ID"].strip()=="DEFAULT" and r["Stage Length"].strip()==str(stage_length)]
    if not selected: raise ValueError(f"no default fixed-point profile for {anp_id}/{operation}/stage {stage_length}")
    selected.sort(key=lambda r:int(r["Point Number"]))
    points=[{"s_m":float(r["Distance (ft)"])*FT_TO_M,"z_m":float(r["Altitude AFE (ft)"])*FT_TO_M,
             "speed_mps":float(r["TAS (kt)"])*KNOT_TO_MPS,"power":float(r["Power Setting"]),
             "bank_deg":0.0,"source_point":int(r["Point Number"])} for r in selected]
    wrows=[r for r in weights if r["ACFT_ID"].strip()==anp_id and r["Stage Length"].strip()==str(stage_length)]
    if len(wrows)!=1: raise ValueError(f"expected one default weight for {anp_id}/stage {stage_length}")
    result={"anp_id_internal":anp_id,"description_internal":ac["Description"],"engine_type":ac["Engine Type"].strip(),
            "engine_count":int(ac["Number Of Engines"]),"installation_identifier":ac["Lateral Directivity Identifier"].strip(),
            "power_parameter":ac["Power Parameter"].strip(),"power_unit":unit,"weight_lb":float(wrows[0]["Weight (lb)"].replace(",","")),
            "operation":operation,"profile_id":"DEFAULT","stage_length":stage_length,"profile_points":len(points)}
    return points,result


def _profile_index(archive: Path) -> dict:
    """In-memory counts for fixed points, procedures, and weight availability."""
    with ZipFile(archive) as z:
        aircraft=_csv_rows(z,"ANP2.3_Aircraft.csv")
        fixed=_csv_rows(z,"ANP2.3_Default_fixed_point_profiles.csv")
        approach=_csv_rows(z,"ANP2.3_Default_approach_procedural_steps.csv")
        departure=_csv_rows(z,"ANP2.3_Default_departure_procedural_steps.csv")
        weights=_csv_rows(z,"ANP2.3_Default_weights.csv")
        npd=_csv_rows(z,"ANP2.3_NPD_data.csv")
    ac={r["ACFT_ID"].strip():r for r in aircraft}
    fixed_count={}
    for r in fixed:
        k=(r["ACFT_ID"].strip(),r["Op Type"].strip(),r["Profile_ID"].strip(),r["Stage Length"].strip())
        fixed_count[k]=fixed_count.get(k,0)+1
    procedure_count={}
    for op,rows in (("A",approach),("D",departure)):
        for r in rows:
            stage=r.get("Stage Length","").strip() or "1"
            k=(r["ACFT_ID"].strip(),op,r["Profile_ID"].strip(),stage)
            procedure_count[k]=procedure_count.get(k,0)+1
    weight_count={}
    for r in weights:
        k=(r["ACFT_ID"].strip(),r["Stage Length"].strip())
        weight_count[k]=weight_count.get(k,0)+1
    npd_modes={(r["NPD_ID"].strip(),r["Noise Metric"].strip(),r["Op Mode"].strip()) for r in npd}
    return {"aircraft":ac,"fixed":fixed_count,"procedural":procedure_count,"weights":weight_count,"npd_modes":npd_modes}


def _official_substitutions(path: Path) -> dict[str,list[dict]]:
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        rows=list(wb["by ICAO code"].iter_rows(values_only=True))
        headers=[str(x).strip() if x is not None else "" for x in rows[0]]
        result={}
        for values in rows[1:]:
            row=dict(zip(headers,values)); code=str(row.get("ICAO_CODE") or "").strip().upper()
            if code: result.setdefault(code,[]).append(row)
        return result
    finally:
        wb.close()


def _mapping_for(code: str, operation: str, index: dict, substitutions: dict) -> dict:
    op=OPERATION_CODE[operation]
    if code in MANUAL_SCENARIOS:
        m=MANUAL_SCENARIOS[code]
        prof=index["fixed"].get((m["anp_id"],op,"DEFAULT","1"),0)
        if not prof:
            return {"status":"manual_scenario_has_no_fixed_profile","mapping_kind":m["mapping_kind"]}
        if code=="P28A" and operation=="arrival":
            return {"status":"scenario_profile_but_operation_excluded_for_NPD_range","mapping_kind":m["mapping_kind"],
                    "anp_id_internal":m["anp_id"],"fixed_profile_points":prof,
                    "reason":"PA28 final landing-roll profile power 1000–1167 rpm/engine is below 1500 rpm NPD minimum and exceeds Doc 29's 5 dB extrapolation caution; exclude the entire arrival event."}
        return {"status":"scenario_fixed_profile_available","mapping_kind":m["mapping_kind"],
                "anp_id_internal":m["anp_id"],"delta_db":m["delta_db"],"official_substitution_used":False}
    # Exact identifier text alone is not enough to claim an aircraft identity.
    if code in index["aircraft"]:
        ac=index["aircraft"][code]
        if index["fixed"].get((code,op,"DEFAULT","1"),0):
            return {"status":"label_match_requires_aircraft_identity_qualification","anp_id_internal":code,
                    "aircraft_description_internal":ac["Description"]}
    rows=substitutions.get(code,[])
    if not rows:
        return {"status":"no_official_2018_ICAO_substitution_row"}
    delta_col="DELTA_APP_dB" if op=="A" else "DELTA_DEP_dB"
    usable=[]
    for row in rows:
        try: delta=float(row[delta_col])
        except (KeyError,TypeError,ValueError): continue
        proxy=str(row.get("ANP_PROXY") or "").strip()
        usable.append((delta,proxy,row))
    if not usable: return {"status":"official_substitution_has_no_operation_delta","official_row_count":len(rows)}
    delta,proxy,row=max(usable,key=lambda x:(x[0],x[1]))
    if proxy not in index["aircraft"]:
        return {"status":"official_max_delta_proxy_absent_from_anp","official_row_count":len(rows),
                "max_delta_db":delta,"selected_proxy_internal":proxy}
    fixed=index["fixed"].get((proxy,op,"DEFAULT","1"),0)
    procedural=index["procedural"].get((proxy,op,"DEFAULT","1"),0)
    w=index["weights"].get((proxy,"1"),0)
    ac=index["aircraft"][proxy]
    npd_id=ac["NPD_ID"].strip()
    npd_ready=(npd_id,"SEL",op) in index["npd_modes"] and (npd_id,"LAmax",op) in index["npd_modes"]
    if not fixed:
        status="procedural_only_not_fixed_profile" if procedural and w and npd_ready else "official_proxy_missing_fixed_profile"
        return {"status":status,"official_row_count":len(rows),"anp_proxy_internal":proxy,"max_delta_db":delta,
                "variant_internal":str(row.get("AIRCRAFT_VARIANT") or ""),"procedure_steps":procedural,
                "weight_available":bool(w),"SEL_and_LAmax_NPD_available":npd_ready,
                "profile_use":"procedure rows require Appendix B synthesis; never coerced to fixed-point geometry"}
    if not (w and npd_ready):
        return {"status":"official_proxy_missing_weight_or_noise_metric","official_row_count":len(rows),"anp_proxy_internal":proxy,"max_delta_db":delta}
    # Doc 29 Appendix G: choose the largest operation-specific Δ across all
    # published variants. Do not combine it with N movement adjustment.
    return {"status":"official_fixed_profile_available","mapping_kind":"official_2018_ICAO_substitution",
            "official_row_count":len(rows),"anp_id_internal":proxy,"delta_db":delta,
            "variant_internal":str(row.get("AIRCRAFT_VARIANT") or ""),"N_adjustment_applied":False,
            "fixed_profile_points":fixed,"npd_id_internal":npd_id}


def _event_sel(event: dict, profile: list[dict], model_info: dict, sel_table, max_table,
               receiver: tuple[float,float,float], operation: str, delta_db: float) -> dict:
    output=[]
    for seg in profile:
        g=_segment_geometry(seg,receiver,operation)
        d=float(g["perpendicular_distance_m"]); p=float(g["power_lb_per_engine"]); v=float(g["speed_mps"])
        if v<=0: raise Doc29InputError("nonpositive profile speed")
        base=npd_level(sel_table,aircraft_id=sel_table.aircraft_id,npd_id=sel_table.npd_id,metric="SEL",operation=operation,
                       power=p,power_unit=sel_table.power_unit,distance_m=d,allow_power_extrapolation=True,minimum_distance_m=30.0)
        lmax=npd_level(max_table,aircraft_id=max_table.aircraft_id,npd_id=max_table.npd_id,metric="LAmax",operation=operation,
                       power=p,power_unit=max_table.power_unit,distance_m=d,allow_power_extrapolation=True,minimum_distance_m=30.0)
        lateral=float(g["lateral_displacement_m"]); beta=float(g["beta_deg"])
        lateral_db=_lateral_attenuation(beta,lateral)
        duration=10.0*math.log10(REFERENCE_SPEED_MPS/v)
        q_ft=float(g["q_m"])/FT_TO_M; length_ft=float(g["length_m"])/FT_TO_M
        ground=seg["start"]["z_m"]<=1.0 and seg["end"]["z_m"]<=1.0
        if ground and operation=="D" and q_ft<=1e-7:
            dlam=(2.0/math.pi)*270.05*10.0**((base-lmax)/10.0); a=length_ft/dlam
            frac=(a/(1+a*a)+math.atan(a))/math.pi
            if frac<=0: raise Doc29InputError("invalid takeoff ground-roll energy fraction")
            finite=10*math.log10(frac); method="Doc29_4-21a_takeoff_ground_roll"
        elif ground and operation=="A" and q_ft>=length_ft-1e-7:
            dlam=(2.0/math.pi)*270.05*10.0**((base-lmax)/10.0); a=-length_ft/dlam
            frac=(-a/(1+a*a)-math.atan(a))/math.pi
            if frac<=0: raise Doc29InputError("invalid landing ground-roll energy fraction")
            finite=10*math.log10(frac); method="Doc29_4-21b_landing_ground_roll"
        else:
            finite=_finite_segment_correction(base,lmax,q_ft,length_ft); method="Doc29_4-20_generic_finite_segment"
        installation=_installation_correction(float(g["phi_deg"]),model_info["installation_identifier"])
        sor=0.0
        if ground and operation=="D" and q_ft<0:
            if model_info["engine_type"]=="Jet": sor,_,_=_start_of_roll_jet(q_ft,float(g["cross_track_m"])/FT_TO_M)
            elif model_info["engine_type"]=="Turboprop": sor,_,_=_start_of_roll_turboprop(q_ft,float(g["cross_track_m"])/FT_TO_M)
        level=base+_impedance_adjustment(15.0,760.0)+duration+installation-lateral_db+finite+sor+delta_db
        output.append({"segment_id":seg["segment_id"],"segment_sel_db_internal":level,"power_unit":sel_table.power_unit,
                       "power_per_engine_internal":p,"speed_mps":v,"distance_m":d,"delta_db_applied_once_to_event_level":delta_db,
                       "installation_correction_db":installation,"lateral_attenuation_db":lateral_db,
                       "finite_segment_correction_db":finite,"finite_segment_method":method,"start_of_roll_correction_db":sor})
    return {"candidate_event_index_internal":event["event_index_internal"],"event_time_utc":event["event_time_utc"],
            "operation":event["operation"],"runway_end_candidates":event["runway_end_candidates"],
            "event_sel_db_internal":energy_sum_db(r["segment_sel_db_internal"] for r in output),
            "segment_count":len(output),"segment_results_internal":output}


def run(events_path: Path, archive: Path, substitutions_path: Path, type_review_path: Path, sites_path: Path,
        observations_path: Path, runway_ends_path: Path) -> dict:
    events_doc=json.loads(events_path.read_text()); runway_doc=json.loads(runway_ends_path.read_text())
    type_review=json.loads(type_review_path.read_text())
    reviewed_candidates={r["icao_type_code"]:r for r in type_review["type_codes"] if r.get("anp_candidate_id_internal_only")}
    index=_profile_index(archive); substitutions=_official_substitutions(substitutions_path)
    all_events=events_doc["events_internal_only"]
    operation_counts={}
    for operation,op in OPERATION_CODE.items():
        type_counts={}
        for e in all_events:
            if e["operation"]==operation: type_counts[e["type_code"]]=type_counts.get(e["type_code"],0)+1
        operation_counts[operation]=type_counts
    type_operation_inventory=[]; mappings={}
    for operation,counts in operation_counts.items():
        for code,count in sorted(counts.items()):
            mapping=_mapping_for(code,operation,index,substitutions)
            candidate=reviewed_candidates.get(code)
            if candidate:
                proxy=candidate["anp_candidate_id_internal_only"]
                ac=index["aircraft"].get(proxy)
                if ac:
                    op=OPERATION_CODE[operation]; npd_id=ac["NPD_ID"].strip()
                    fixed=index["fixed"].get((proxy,op,"DEFAULT","1"),0)
                    procedural=index["procedural"].get((proxy,op,"DEFAULT","1"),0)
                    mapping["unapproved_family_proxy_candidate_internal_only"]={
                        "anp_id_internal":proxy,"review_status":candidate["mapping_status"],
                        "review_rationale":candidate.get("candidate_rationale"),"fixed_point_rows":fixed,
                        "default_procedural_rows":procedural,"default_stage1_weight_available":bool(index["weights"].get((proxy,"1"),0)),
                        "SEL_and_LAmax_available":(npd_id,"SEL",op) in index["npd_modes"] and (npd_id,"LAmax",op) in index["npd_modes"],
                        "profile_class":"fixed_point" if fixed else ("procedural_only" if procedural else "no_default_profile"),
                        "accepted_for_calculation":False}
            mappings[(code,operation)]=mapping
            type_operation_inventory.append({"icao_type_code":code,"operation":operation,"candidate_event_windows_not_confirmed_movements":count,
                                              "mapping_and_fixed_profile_status":mapping})
    sites,obs=_read_sites(sites_path,observations_path)
    runway_ends={str(f["attributes"]["RWY_END_ID"]):f["attributes"] for f in runway_doc["features"] if f["attributes"].get("ARPT_ID")=="VNY"}
    for v in runway_ends.values():
        if v.get("EFF_DATE")!="2026/10/01": raise ValueError("unexpected FAA runway geometry vintage")
    qualified=[]; excluded_quality=0
    for event in all_events:
        operation=event["operation"]
        required=("approach_witness","ground_witness") if operation=="arrival" else ("ground_witness","climb_base_witness","climb_witness")
        if not all(event[k].get("position_source")=="adsb_icao" for k in required):
            excluded_quality+=1; continue
        event=dict(event); event["event_index_internal"]=len(qualified)+1
        event["event_time_utc"]=event["ground_witness"]["timestamp_utc"]
        qualified.append(event)
    qualified_counts={}
    for e in qualified:
        key=(e["type_code"],e["operation"])
        qualified_counts[key]=qualified_counts.get(key,0)+1
    mapping_cache={}; modeled=[]; excluded_model=[]
    for event in qualified:
        code=event["type_code"]; operation=event["operation"]; mapping=mappings[(code,operation)]
        if mapping.get("status") not in ("scenario_fixed_profile_available","official_fixed_profile_available"):
            continue
        anp_id=mapping["anp_id_internal"]; op=OPERATION_CODE[operation]; key=(anp_id,op)
        if key not in mapping_cache:
            raw,info=load_fixed_profile(archive,anp_id,op)
            constructed=segment_profile_points(raw,op=="A")
            profile=[]
            for i,(a,b) in enumerate(zip(constructed,constructed[1:]),1):
                if math.hypot(b["s_m"]-a["s_m"],b["z_m"]-a["z_m"])<=0: raise ValueError("nonpositive fixed-profile segment")
                profile.append({"segment_id":i,"start":a,"end":b})
            match,sel=load_anp_npd_table(archive,anp_id,"SEL",op); maxmatch,lamax=load_anp_npd_table(archive,anp_id,"LAmax",op)
            if sel.power_unit!=info["power_unit"] or lamax.power_unit!=info["power_unit"]: raise ValueError("profile/NPD power-unit mismatch")
            mapping_cache[key]=(raw,info,profile,sel,lamax)
        raw,info,profile,sel,lamax=mapping_cache[key]
        for site in sites:
            if site.get("spatial_modeling_status"): continue
            ends=sorted(set(event["runway_end_candidates"]))
            offsets={}
            for end_id in ends:
                a=runway_ends[end_id]; x,y,extra=_local_track_offsets({"lat":float(a["LAT_DECIMAL"]),"lon":float(a["LONG_DECIMAL"])},site,float(a["TRUE_ALIGNMENT"]))
                offsets[end_id]=(x,y)
            for choice in ("L","R"):
                candidates=ends
                end_id=candidates[0] if len(candidates)==1 else next(x for x in candidates if x.endswith(choice))
                x,y=offsets[end_id]
                for height in RECEIVER_HEIGHTS_M:
                    for vertical_offset in VERTICAL_OFFSETS_M:
                        try:
                            result=_event_sel(event,profile,info,sel,lamax,(x,y,height+vertical_offset),op,mapping.get("delta_db",0.0))
                        except Doc29InputError as exc:
                            excluded_model.append({"type_code":code,"operation":operation,"reason":str(exc),"mapping_kind":mapping["mapping_kind"]})
                            break
                        result.update({"station_id":site["station_id"],"runway_assignment":choice,"runway_end_used":end_id,
                                       "assumed_mic_height_m":height,"illustrative_site_plane_offset_m":vertical_offset,
                                       "mapping_kind":mapping["mapping_kind"],"mapped_anp_variant_internal":anp_id})
                        modeled.append(result)
    # Remove duplicate failure messages from separate receivers and retain counts.
    excluded_counts={}
    for x in excluded_model:
        k=(x["type_code"],x["operation"],x["reason"]); excluded_counts[k]=excluded_counts.get(k,0)+1
    summaries=[]
    for site in sites:
        if site.get("spatial_modeling_status"): continue
        for choice in ("L","R"):
            for height in RECEIVER_HEIGHTS_M:
                for offset in VERTICAL_OFFSETS_M:
                    rows=[r for r in modeled if r["station_id"]==site["station_id"] and r["runway_assignment"]==choice and r["assumed_mic_height_m"]==height and r["illustrative_site_plane_offset_m"]==offset]
                    if rows:
                        bypart={}
                        for r in rows:
                            part=_local_daypart(r["event_time_utc"]); bypart.setdefault(part,[]).append({"event_time_utc":r["event_time_utc"],"event_sel_db_internal":r["event_sel_db_internal"]})
                        # Maintain the actual event SEL rows and apply local CNEL weights.
                        events_cnel=[{"event_time_utc":r["event_time_utc"],"event_sel_db_internal":r["event_sel_db_internal"]} for r in rows]
                        obs_db=site["observed_daily_cnel_context_db"]
                        cnel=_subset_cnel(events_cnel)
                        summaries.append({"station_id":site["station_id"],"runway_assignment":choice,"mic_height_m":height,"site_plane_offset_m":offset,
                                          "unique_event_window_count":len(events_cnel),"subset_event_SEL_sum_db_internal":energy_sum_db(r["event_sel_db_internal"] for r in rows),
                                          "subset_daily_CNEL_db_internal":cnel["subset_event_cnel_db_internal"],"event_count_by_CNEL_daypart":cnel["candidate_event_counts_by_daypart"],
                                          "observed_total_aircraft_CNEL_context_db":obs_db,
                                          "subset_CNEL_minus_observed_total_db_exploratory_only":None if obs_db is None else cnel["subset_event_cnel_db_internal"]-obs_db,
                                          "interpretation":"internal exploratory subset diagnostic; incomplete activity/mapping/profile; not validation, lower bound, or complete modeled exposure"})
    return {"schema":"quiet_la_vny_generic_fixed_profile_event_adapter_v1_internal",
            "date_local":events_doc["date_local"],"activity_source_sha256":events_doc["activity_sha256"],
            "candidate_activity_counts_not_confirmed_movements":events_doc["counts"],
            "qualification":{"events_before_witness_quality_filter":len(all_events),"events_after_required_ADSB_ICAO_witness_filter":len(qualified),
                             "events_excluded_for_witness_quality":excluded_quality,
                             "quality_filtered_windows_by_operation":{op:sum(1 for e in qualified if e["operation"]==op) for op in OPERATION_CODE},
                             "quality_filtered_windows_by_type_and_operation_internal_only":[{"icao_type_code":k[0],"operation":k[1],"candidate_windows":v} for k,v in sorted(qualified_counts.items())],
                             "eligible_type_operation_pairs":len({(x["type_code"],x["operation"]) for x in qualified if mappings[(x["type_code"],x["operation"])].get("status") in ("scenario_fixed_profile_available","official_fixed_profile_available")}),
                             "unique_candidate_event_windows_modeled":len({r["candidate_event_index_internal"] for r in modeled}),
                             "candidate_event_windows_excluded_for_known_NPD_range_issue":sum(1 for e in qualified if mappings[(e["type_code"],e["operation"])].get("status")=="scenario_profile_but_operation_excluded_for_NPD_range"),
                             "modeled_event_receiver_sensitivity_rows":len(modeled),
                             "exclusions_for_NPD_or_physics":[{"type_code":k[0],"operation":k[1],"reason":k[2],"failed_cell_count":v} for k,v in sorted(excluded_counts.items())]},
            "mapping_policy":{"exact_icao_anp_code_string_alone":"not accepted as aircraft identity; requires an independently qualified crosswalk","official_substitution":"choose operation-specific max Δ across table variants; require the selected proxy's fixed profile, SEL/LAmax NPD and weight; never apply N simultaneously","manual_scenarios":MANUAL_SCENARIOS,
                              "profile_policy":"procedural-only aircraft are inventoried but excluded from this fixed-profile adapter pending Doc 29 Appendix B synthesis"},
            "type_operation_inventory_internal_only":type_operation_inventory,
            "fixed_profile_scenario_mapping_counts":{status:sum(x["mapping_and_fixed_profile_status"].get("status")==status for x in type_operation_inventory) for status in sorted({x["mapping_and_fixed_profile_status"].get("status") for x in type_operation_inventory})},
            "mapped_profile_ids_and_properties_internal_only":[{"anp_id_internal":key[0],"operation":key[1],"engine_type":val[1]["engine_type"],"power_unit":val[1]["power_unit"],"installation_identifier":val[1]["installation_identifier"],"fixed_profile_points":val[1]["profile_points"],"segments_after_doc29_subdivision":len(val[2])} for key,val in mapping_cache.items()],
            "monitor_context":{"modeled_sites":sorted({r["station_id"] for r in modeled}),"excluded_sites":[s["station_id"] for s in sites if s.get("spatial_modeling_status")],
                               "height_scenarios_m":list(RECEIVER_HEIGHTS_M),"illustrative_site_plane_offsets_m":list(VERTICAL_OFFSETS_M),"candidate_subset_summaries":summaries,
                               "site_and_mic_datum_uncertainty":"2021 site coordinates only, 0 mic heights documented; plane offsets are illustrative and not a physical uncertainty bound; no terrain/geoid conversion"},
            "model_assumptions":{"flat_soft_ground":True,"terrain_or_building_blockage_modelled":False,"weather":"15 C and 760 mmHg; no hourly weather correction","reference_speed_knots":160,
                                 "aircraft_geometry":"default fixed profile aligned to inferred FAA runway end; ADS-B geometry selects event/runway only, not acoustic path"},
            "activity_rows_internal_only":modeled,
            "sources":{"event_windows_sha256":_sha(events_path),"anp_archive_sha256":_sha(archive),"official_substitution_xlsx_sha256":_sha(substitutions_path),
                       "type_review_sha256":_sha(type_review_path),
                       "site_coordinates_sha256":_sha(sites_path),"daily_CNEL_sha256":_sha(observations_path),"runway_end_capture_sha256":_sha(runway_ends_path),
                       "adapter_code_sha256":_sha(Path(__file__))}}


def main():
    p=argparse.ArgumentParser()
    for name in ("events","anp","substitutions","type-review","sites","observations","runway-ends","output"):
        p.add_argument("--"+name,type=Path,required=True)
    a=p.parse_args(); result=run(a.events,a.anp,a.substitutions,a.type_review,a.sites,a.observations,a.runway_ends)
    a.output.parent.mkdir(parents=True,exist_ok=True); a.output.write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps({"event_windows_total":result["qualification"]["events_before_witness_quality_filter"],
                      "quality_filtered":result["qualification"]["events_after_required_ADSB_ICAO_witness_filter"],
                      "modeled_event_receiver_rows":result["qualification"]["modeled_event_receiver_sensitivity_rows"],
                      "modeled_sites":result["monitor_context"]["modeled_sites"]},indent=2))


if __name__=="__main__": main()
